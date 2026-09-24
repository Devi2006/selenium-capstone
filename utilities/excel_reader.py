from openpyxl import load_workbook


def read_test_data(file_path, sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    headers = [cell.value for cell in sheet[1]]

    row = next(sheet.iter_rows(min_row=2, values_only=True))

    return dict(zip(headers, row))